"""Build two simple assessment sheets and their column guide."""
import hashlib
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from annotate import ROOT, PRIVATE
from review import read_rows, index_rows


COMMON = [('journal', 'Journal', 'SCORE journal title.'),
          ('year', 'Year', 'Indexed publication year; descending within journal.'),
          ('title', 'Article title', 'Full article title.'),
          ('doi', 'DOI', 'Persistent article DOI.'),
          ('authors', 'Authors', 'Bibliographic author list; not supplied for abstract screening.'),
          ('available', 'Full text available locally', 'YES only for a source file with a matching recorded checksum.'),
          ('local_file', 'Local full text', 'Path to the locally available source.'),
          ('version_note', 'Source version note', 'Known manuscript differences or missing supporting assets.')]
ABSTRACT = COMMON + [('annotation', 'Annotation', 'YES / NO / UNCLEAR using only journal, title, abstract and keywords. Blank means not yet assessed.'),
                     ('abstract', 'Abstract', 'Original screening input; private licensed metadata.'),
                     ('keywords', 'Keywords', 'Original screening input; private licensed metadata.')]
FULLTEXT = COMMON + [('annotation', 'Annotation', 'YES / NO / UNCLEAR based on substantive source review. Blank means not yet reviewed.'),
                     ('evidence', 'Evidence', 'Short source excerpts with page or section references; links identify supporting sources.'),
                     ('reasoning', 'Reasoning', 'Why the same participant responses do or do not satisfy both eligibility criteria, and any unresolved source issues.')]


def format_sheet(sheet, rows, columns, table_name):
    sheet.append([name for _, name, _ in columns])
    for row in rows:
        sheet.append([row.get(key, '') for key, _, _ in columns])
    sheet.freeze_panes = 'C2'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='264E58')
    for number, (key, _, _) in enumerate(columns, 1):
        width = 10 if key == 'year' else 18 if key in ['annotation', 'available'] else 58 if key in ['title', 'evidence', 'reasoning', 'abstract', 'keywords'] else 32
        sheet.column_dimensions[get_column_letter(number)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 55
    table = Table(displayName=table_name, ref=f'A1:{get_column_letter(len(columns))}{sheet.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    sheet.add_table(table)
    for number, record in enumerate(rows, 2):
        if record.get('doi'):
            sheet.cell(number, 4).hyperlink = 'https://doi.org/' + record['doi']
        if record.get('local_file'):
            sheet.cell(number, 7).hyperlink = Path(record['local_file']).as_uri()


def main():
    articles = sorted(read_rows(PRIVATE / 'articles.csv'), key=lambda r: (r['journal'], -int(r['year']), r['title']))
    access = index_rows(read_rows(PRIVATE / 'fulltext.csv'))
    predictions = index_rows(read_rows(PRIVATE / 'predictions.csv'))
    references = index_rows(read_rows(PRIVATE / 'fulltext_reviews.csv'))
    abstract_rows, fulltext_rows = [], []
    for article in articles:
        key = article['article_id']
        source = access.get(key, {})
        path = Path(source.get('fulltext_path', ''))
        verified = path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest() == source.get('source_sha256')
        base = dict(article, available='YES' if verified else 'NO',
                    local_file=str(path) if verified else '', version_note=source.get('version_note', ''))
        abstract_rows.append(dict(base, annotation=predictions.get(key, {}).get('annotation', '')))
        fulltext_rows.append(dict(base, **{field: references.get(key, {}).get(field, '') for field in ['annotation', 'evidence', 'reasoning']}))
    workbook = Workbook()
    workbook.remove(workbook.active)
    format_sheet(workbook.create_sheet('Abstract screening'), abstract_rows, ABSTRACT, 'AbstractScreening')
    format_sheet(workbook.create_sheet('Full-text review'), fulltext_rows, FULLTEXT, 'FullTextReview')
    guide = workbook.create_sheet('Column guide')
    guide.append(['Sheet', 'Variable', 'Meaning'])
    for name, columns in [('Abstract screening', ABSTRACT), ('Full-text review', FULLTEXT)]:
        for _, title, description in columns:
            guide.append([name, title, description])
    guide.append(['Both', 'Review status', 'AI assessments are provisional; no human verification has been completed. Source downloads and source reviews are distinct.'])
    guide.append(['Both', 'Prompt', 'The sole screening prompt is score/prompt.md; full-text instructions adapt input and output only.'])
    guide.column_dimensions['A'].width = 24
    guide.column_dimensions['B'].width = 32
    guide.column_dimensions['C'].width = 110
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    workbook.save(PRIVATE / 'SCORE_validation_620.xlsx')
    lines = ['# Spreadsheet codebook', '', 'The workbook has two 620-row assessment sheets, each with one annotation column, sorted by journal and descending year. Abstract screening and full-text review remain separate. No article ID is displayed; code joins records internally.', '',
             'Labels: YES = qualifies; NO = clearly fails; UNCLEAR = unresolved. Blank means not yet assessed. Full-text availability does not imply review. Evidence and reasoning accompany only the full-text assessment. Full-text AI judgments remain provisional pending human verification.', '',
             '| Sheet | Variable | Meaning |', '| --- | --- | --- |']
    for name, columns in [('Abstract screening', ABSTRACT), ('Full-text review', FULLTEXT)]:
        for _, title, description in columns:
            lines.append(f'| {name} | {title} | {description} |')
    (ROOT / 'score/codebook.md').write_text('\n'.join(lines) + '\n')
    print(f'Built workbook: {len(articles)} rows in each assessment sheet; one annotation column per sheet.')


if __name__ == '__main__':
    main()
